# -*- coding: utf-8 -*-
"""
한양이엔지 통합 관제 프로그램 (Monitor + Viewer)
=================================================
하나의 독립 실행 프로그램으로 다음 두 기능을 동시에 수행합니다.

1. [백그라운드 모니터링 — 구 Client.py]
   이 PC에서 S5D / DDWORKS 실행 현황을 감지해 관제 서버로 주기적으로 전송합니다.

2. [관제 대시보드 — 구 viewer_v4.py]
   전체 구역의 실시간 접속 현황을 차트/명단으로 보여줍니다.

· 시스템 트레이에 상주하며, 창을 닫아도 백그라운드 모니터링은 계속됩니다.
· 중복 실행 방지(단일 인스턴스) 및 Windows 부팅 시 자동 실행(트레이 모드)을 지원합니다.
"""

import tkinter as tk
from tkinter import ttk
import csv
import json
import time
import socket
import threading
import logging
import requests
import xlwings as xw
import os
import re
import sys
import matplotlib
matplotlib.use("TkAgg")
import matplotlib.pyplot as plt
from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg
import matplotlib.font_manager as fm
from tkinter import filedialog
from datetime import datetime
from tkinter import scrolledtext

# 💡 백그라운드 모니터링(구 Client.py)에 필요한 모듈
import psutil
import ctypes
from PIL import Image, ImageDraw
import pystray

# 💡 Windows 전용 모듈은 플랫폼 가드 (다른 OS에서도 import만은 가능하도록)
IS_WINDOWS = (os.name == 'nt')
if IS_WINDOWS:
    import winreg

# ==========================================
# 🔧 [설정] 서버 및 프로그램 정보
# ==========================================
SERVER_URL = "http://12.26.204.100:5000"
API_KEY = "HanyangENG-Monitor-2026!"

# ==========================================
# 🔧 [설정] 백그라운드 모니터링 (구 Client.py)
# ==========================================
# 💡 모니터링 결과를 전송할 서버 엔드포인트
STATUS_REPORT_URL = f"{SERVER_URL}/receive_status"
# 💡 감지 대상 프로세스 (S5D, DDWORKS)
TARGET_PROCESSES = ["S5D.exe", "dinno.hu3d.wpf.hookupdesigner.exe"]
# 💡 상태 전송 주기 (초)
CHECK_INTERVAL = 10
# 💡 Windows 시작프로그램 등록 이름
AUTOSTART_NAME = "HanyangENG_Monitor"
# 💡 단일 인스턴스 보장용 뮤텍스 이름
MUTEX_NAME = "HanyangENG_Monitor_Unified_Mutex"
# 💡 S5D 시작 후 창 제목 조회 유예시간 (초) — Tiara 보안 모듈 초기화 회피
S5D_STARTUP_GRACE_SEC = 15
# 💡 트레이 모드(부팅 자동 실행)로 시작할 때의 커맨드라인 인자
TRAY_START_FLAG = "--minimized"

# 💡 프로세스 파일명 → 표시 이름
# ─ 클라이언트 TARGET_PROCESSES 와 동일하게 맞춰주세요
PROGRAMS_MAP = {
    "s5d.exe": "S5D",
    "dinno.hu3d.wpf.hookupdesigner.exe": "DDWORKS"
}

# 💡 구역명 화면 표시 변환 테이블
# ─ 프로세스 창제목에서 추출되는 내부 이름(key) → 화면에 보여줄 이름(value)
# ─ 여기에 없는 구역은 자동으로 언더바(_)를 하이픈(-)으로 바꿔서 표시합니다.
# ─ 추후 변환이 필요하면 아래에 한 줄만 추가하면 됩니다!
ZONE_DISPLAY_MAP = {
    "LSI": "S3",
    "P4_UW": "P4-UW",
    "P4_UE": "P4-UE",
    # 💡 프로세스별 분리 표시 (구역:프로세스명 → 화면 표시명)
    "P4_I:s5d.exe": "P4-I (S5D)",
    "P4_I:dinno.hu3d.wpf.hookupdesigner.exe": "P4-I (DDWORKS)",
    # "예시_내부명": "예시 표시명",  ← 이런 식으로 추가
}

# 💡 지역 버튼 표시명 → 내부 지역 키
REGION_BUTTONS = {
    "평택": "평택",
    "기흥•화성": "화성",
    "천안": "천안",
    "해외": "해외"
}

# 💡 지역별 구역 설정
# ─ "grouped": True → 그룹(P1,P2...) + 하위 라인 계층형 막대그래프
# ─ "grouped": False → 라인만 평면 막대그래프
REGION_CONFIG = {
    "평택": {
        "grouped": True,
        "groups": {
            "P1": ["P1_1", "P1_2", "P1_3"],
            "P2": ["P2_1", "P2_2", "P2_3", "P2_4"],
            "P3": ["P3_1", "P3_2", "P3_3", "P3_4", "P3_I"],
            "P4": ["P4_1", "P4_I:s5d.exe", "P4_I:dinno.hu3d.wpf.hookupdesigner.exe", "P4_UW", "P4_UE"],
        }
    },
    "화성": {
        "grouped": False,
        "zones": ["V1", "LSI", "17L"]
    },
    "천안": {
        "grouped": False,
        "zones": ["C"]
    },
    "해외": {
        "grouped": False,
        "zones": ["T", "X2"]
    }
}

# 💡 PyInstaller .exe 배포 시에도 설정 파일이 .exe와 같은 폴더에 생성되도록 경로 설정
def _get_app_dir():
    """실행 파일(.exe) 또는 스크립트(.py)가 있는 폴더 경로 반환"""
    if getattr(sys, 'frozen', False):
        # PyInstaller로 빌드된 .exe 실행 시
        return os.path.dirname(sys.executable)
    else:
        # 일반 Python 스크립트 실행 시
        return os.path.dirname(os.path.abspath(__file__))

APP_DIR = _get_app_dir()
SETTINGS_FILE = os.path.join(APP_DIR, "viewer_settings.json")
# 💡 백그라운드 모니터링 디버그 로그 (.exe와 같은 폴더)
LOG_FILE = os.path.join(APP_DIR, "monitor_debug.log")
# 💡 서버에서 내려받는 임시 파일도 쓰기 가능한 APP_DIR에 생성
TEMP_EMP_FILE = os.path.join(APP_DIR, "temp_emp.xlsx")
TEMP_LOG_FILE = os.path.join(APP_DIR, "temp_logs.csv")

def get_display_name(zone_name):
    """내부 구역명 → 화면 표시명 변환
    1순위: ZONE_DISPLAY_MAP에 등록된 이름
    2순위: 언더바(_)를 하이픈(-)으로 자동 변환
    """
    if zone_name in ZONE_DISPLAY_MAP:
        return ZONE_DISPLAY_MAP[zone_name]
    return zone_name.replace("_", "-")

def load_settings():
    try:
        with open(SETTINGS_FILE, 'r', encoding='utf-8') as f:
            data = json.load(f)
            if "favorite_zones" not in data:
                data["favorite_zones"] = ["T", "C", "P4_1", "P3_4"]
            if "refresh_interval" not in data:
                data["refresh_interval"] = "수동"
            return data
    except:
        return {
            "default_region": "__favorites__",
            "favorite_zones": ["T", "C", "P4_1", "P3_4"],
            "refresh_interval": "수동",
        }

def save_settings(settings):
    with open(SETTINGS_FILE, 'w', encoding='utf-8') as f:
        json.dump(settings, f, ensure_ascii=False, indent=2)


# ==========================================================
# 🖥 백그라운드 모니터링 에이전트 (구 Client.py)
#   이 PC의 S5D / DDWORKS 실행 현황을 감지해 서버로 전송한다.
# ==========================================================
class MonitoringAgent:
    def __init__(self, log_callback=None):
        self.stop_event = threading.Event()
        self.failed_queue = []
        self._process_first_seen = {}     # {proc_name: 최초 감지 timestamp}
        self._log_callback = log_callback  # 뷰어 디버그 로그로 전달할 콜백
        self.thread = None

        # 💡 GUI 표시용 상태 요약
        self.summary = "준비 중…"
        self.summary_color = "#555555"

    # ── 로그 (콘솔 + 파일 + 뷰어 디버그 창) ──
    def log(self, msg):
        timestamp = datetime.now().strftime("%H:%M:%S")
        line = f"[{timestamp}] [모니터] {msg}"
        print(line)
        try:
            with open(LOG_FILE, 'a', encoding='utf-8') as f:
                f.write(line + "\n")
        except Exception:
            pass
        if self._log_callback:
            try:
                self._log_callback(line)
            except Exception:
                pass

    # ── 네트워크 / 프로세스 ──
    def get_local_ip(self):
        try:
            s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
            s.connect(("8.8.8.8", 80))
            ip = s.getsockname()[0]
            s.close()
            return ip
        except Exception:
            return socket.gethostbyname(socket.gethostname())

    def is_process_running(self, process_name):
        """프로세스 실행 여부만 가볍게 확인 (창 제목은 읽지 않음)"""
        for proc in psutil.process_iter(['name']):
            try:
                if proc.info['name'] and proc.info['name'].lower() == process_name.lower():
                    return True
            except (psutil.NoSuchProcess, psutil.AccessDenied):
                pass
        return False

    def get_detailed_window_title(self, process_name):
        """프로세스의 보이는 창 제목들을 ' / '로 이어 반환 (Windows 전용)"""
        if not IS_WINDOWS:
            return None

        target_pids = []
        for proc in psutil.process_iter(['name', 'pid']):
            try:
                if proc.info['name'] and proc.info['name'].lower() == process_name.lower():
                    target_pids.append(proc.info['pid'])
            except (psutil.NoSuchProcess, psutil.AccessDenied):
                pass

        if not target_pids:
            return None

        EnumWindows = ctypes.windll.user32.EnumWindows
        EnumWindowsProc = ctypes.WINFUNCTYPE(ctypes.c_bool, ctypes.POINTER(ctypes.c_int), ctypes.POINTER(ctypes.c_int))
        GetWindowThreadProcessId = ctypes.windll.user32.GetWindowThreadProcessId
        GetWindowText = ctypes.windll.user32.GetWindowTextW
        GetWindowTextLength = ctypes.windll.user32.GetWindowTextLengthW
        IsWindowVisible = ctypes.windll.user32.IsWindowVisible

        found_titles = []

        def foreach_window(hwnd, lParam):
            if IsWindowVisible(hwnd):
                pid = ctypes.c_ulong()
                GetWindowThreadProcessId(hwnd, ctypes.byref(pid))
                if pid.value in target_pids:
                    length = GetWindowTextLength(hwnd)
                    if length > 0:
                        buff = ctypes.create_unicode_buffer(length + 1)
                        GetWindowText(hwnd, buff, length + 1)
                        found_titles.append(buff.value)
            return True

        EnumWindows(EnumWindowsProc(foreach_window), 0)
        return " / ".join(found_titles) if found_titles else "창 제목 없음 (백그라운드)"

    # ── 상태 수집 + 서버 전송 ──
    def send_status(self):
        hostname = socket.gethostname()
        ip_address = self.get_local_ip()
        now_ts = time.time()
        running_map = {}

        for proc_name in TARGET_PROCESSES:
            # 1단계: 실행 여부만 먼저 확인 (가벼운 체크)
            running = self.is_process_running(proc_name)
            running_map[proc_name] = running

            if not running:
                if proc_name in self._process_first_seen:
                    del self._process_first_seen[proc_name]
                    self.log(f"{proc_name} 종료 감지")
                self.failed_queue.append({
                    "ip": ip_address, "hostname": hostname,
                    "process_name": proc_name, "is_running": False, "details": ""
                })
                continue

            # 실행 중 → 최초 감지 시각 기록
            if proc_name not in self._process_first_seen:
                self._process_first_seen[proc_name] = now_ts
                self.log(f"{proc_name} 새로 감지 — {S5D_STARTUP_GRACE_SEC}초 유예 후 창 제목 조회")

            # 2단계: S5D는 시작 후 유예시간 동안 창 제목 조회 스킵 (Tiara 보안 모듈 회피)
            elapsed = now_ts - self._process_first_seen[proc_name]
            if proc_name.lower() == "s5d.exe" and elapsed < S5D_STARTUP_GRACE_SEC:
                self.failed_queue.append({
                    "ip": ip_address, "hostname": hostname,
                    "process_name": proc_name, "is_running": True,
                    "details": f"초기화 대기 중 ({int(S5D_STARTUP_GRACE_SEC - elapsed)}초)"
                })
                continue

            # 3단계: 유예 완료 → 창 제목 조회
            window_details = self.get_detailed_window_title(proc_name)

            # 4단계: DDWORKS는 Hookup Designer 창([구역명])이 열려야 유효
            if proc_name.lower() == "dinno.hu3d.wpf.hookupdesigner.exe":
                if not window_details or "[" not in window_details:
                    self.failed_queue.append({
                        "ip": ip_address, "hostname": hostname,
                        "process_name": proc_name, "is_running": False, "details": ""
                    })
                    continue

            self.failed_queue.append({
                "ip": ip_address, "hostname": hostname,
                "process_name": proc_name, "is_running": True,
                "details": window_details if window_details else "창 제목 없음"
            })

        # ── 큐 전송 (실패분은 다시 큐에 남김, 최대 50건) ──
        remaining_queue = []
        send_failed = False
        headers = {"X-API-Key": API_KEY}

        for p in self.failed_queue:
            try:
                res = requests.post(STATUS_REPORT_URL, json=p, headers=headers,
                                    timeout=5, proxies={"http": None, "https": None})
                if res.status_code == 403:
                    self.log("서버에서 미인가 거부 (403) — 데이터 폐기")
                elif res.status_code != 200:
                    remaining_queue.append(p)
                    send_failed = True
            except requests.exceptions.RequestException:
                remaining_queue.append(p)
                send_failed = True

        self.failed_queue = remaining_queue[-50:]
        self._update_summary(running_map, send_failed)

    def _update_summary(self, running_map, send_failed):
        """GUI 상태바에 표시할 요약 문자열 갱신"""
        parts = []
        for proc_name in TARGET_PROCESSES:
            disp = PROGRAMS_MAP.get(proc_name.lower(), proc_name)
            mark = "●" if running_map.get(proc_name) else "○"
            parts.append(f"{disp} {mark}")
        status_str = "  ".join(parts)
        if send_failed:
            self.summary = f"{status_str}   · 서버 전송 지연 (재시도 중)"
            self.summary_color = "#F44336"
        else:
            self.summary = f"{status_str}   · 전송 정상"
            self.summary_color = "#2E7D32"

    # ── 스레드 루프 ──
    def _loop(self):
        self.log("모니터링 스레드 시작")
        while not self.stop_event.is_set():
            try:
                self.send_status()
            except Exception as e:
                self.log(f"send_status 오류: {e}")
            for _ in range(CHECK_INTERVAL):
                if self.stop_event.is_set():
                    break
                time.sleep(1)
        self.log("모니터링 스레드 종료")

    def start(self):
        if self.thread and self.thread.is_alive():
            return
        self.stop_event.clear()
        self.thread = threading.Thread(target=self._loop, daemon=True)
        self.thread.start()

    def stop(self):
        self.stop_event.set()


# ==========================================================
# 🔧 단일 인스턴스 / 자동 실행 / 트레이 (구 Client.py)
# ==========================================================
_mutex_handle = None

def acquire_single_instance():
    """중복 실행 방지. 이미 실행 중이면 안내 후 False 반환."""
    global _mutex_handle
    if not IS_WINDOWS:
        return True  # 비 Windows에서는 단일 인스턴스 검사 생략
    try:
        _mutex_handle = ctypes.windll.kernel32.CreateMutexW(None, False, MUTEX_NAME)
        if ctypes.windll.kernel32.GetLastError() == 183:  # ERROR_ALREADY_EXISTS
            ctypes.windll.user32.MessageBoxW(
                0,
                "관제 프로그램이 이미 실행 중입니다!\n(우측 하단 숨겨진 아이콘 영역을 확인하세요.)",
                "실행 안내", 0x30  # MB_ICONWARNING
            )
            return False
    except Exception:
        pass
    return True

def get_executable_path():
    if getattr(sys, 'frozen', False):
        return f'"{sys.executable}"'
    return f'"{sys.executable}" "{os.path.abspath(__file__)}"'

def register_autostart():
    """Windows 부팅 시 트레이 모드로 자동 실행되도록 등록"""
    if not IS_WINDOWS:
        return
    try:
        exe_path = f"{get_executable_path()} {TRAY_START_FLAG}"
        reg_key = winreg.OpenKey(
            winreg.HKEY_CURRENT_USER,
            r"Software\Microsoft\Windows\CurrentVersion\Run",
            0, winreg.KEY_READ | winreg.KEY_WRITE
        )
        try:
            current_val, _ = winreg.QueryValueEx(reg_key, AUTOSTART_NAME)
            if current_val == exe_path:
                winreg.CloseKey(reg_key)
                return
        except FileNotFoundError:
            pass
        winreg.SetValueEx(reg_key, AUTOSTART_NAME, 0, winreg.REG_SZ, exe_path)
        winreg.CloseKey(reg_key)
    except Exception as e:
        print(f"[자동 시작 등록 실패 — 무시됨] {e}")

def create_tray_image():
    image = Image.new('RGBA', (64, 64), color=(0, 0, 0, 0))
    draw = ImageDraw.Draw(image)
    draw.ellipse((4, 4, 60, 60), fill="#00479A", outline="white", width=3)
    return image

def build_tray_icon(viewer):
    """뷰어와 연결된 시스템 트레이 아이콘 생성.
    모든 콜백은 root.after(0, ...)로 tkinter 메인 스레드에서 실행한다."""
    def on_open(icon, item):
        viewer.root.after(0, viewer.show_window)

    def on_refresh(icon, item):
        viewer.root.after(0, viewer.manual_refresh)

    def on_debug(icon, item):
        viewer.root.after(0, viewer.open_debug_window)

    def on_quit(icon, item):
        viewer.root.after(0, viewer.real_quit)

    menu = pystray.Menu(
        pystray.MenuItem('📊 관제 화면 열기', on_open, default=True),
        pystray.MenuItem('🔄 지금 갱신', on_refresh),
        pystray.MenuItem('🔍 디버그 로그', on_debug),
        pystray.Menu.SEPARATOR,
        pystray.MenuItem('❌ 완전 종료', on_quit),
    )
    return pystray.Icon("HanyangMonitor", create_tray_image(), "한양이엔지 통합 관제", menu)

# ==========================================

class MonitorViewer:
    def __init__(self, root):
        self.root = root
        self.root.title("한양이엔지 관제 시스템")
        self.root.geometry("950x700")
        self.root.minsize(950, 700)  # 💡 최소 창크기

        self.colors = {
            "primary": "#00479A",
            "primary_light": "#5C9BD5",
            "secondary": "#E0F7FA",
            "background": "#F5F7FA",
            "card_bg": "#FFFFFF",
            "text": "#212121",
            "header_text": "#FFFFFF",
            "status_online": "#4CAF50",
            "status_offline": "#F44336",
            "btn_normal": "#E8EAF6",
            "btn_active": "#00479A",
            "btn_text_normal": "#333333",
            "btn_text_active": "#FFFFFF",
            "group_bar": "#003C82",
            "sub_bar": "#5C9BD5",
            "empty_bar": "#D0D0D0",
            "over_limit": "#F44336",      # 💡 제한 인원 초과 시 빨간색
            "over_limit_group": "#B71C1C", # 💡 그룹 합계 초과 시 진한 빨간색
        }

        self.root.configure(bg=self.colors["background"])
        self.emp_by_ip = {}
        self.emp_by_pc = {}
        self.emp_info = {}            # 💡 사원명 → {"team": 팀, "gongjong": 공종}
        self.active_users = {}
        self.online_pcs = set()           # 💡 클라이언트가 실행 중인 PC의 IP/PC명 (미실행 포함)
        self.server_zone_limits = {}  # 💡 서버에서 받아온 접속 제한 인원
        self.selected_region = None
        self.region_buttons = {}
        self.clicked_zone = None          # 클릭한 개별 구역 (내부명)
        self.clicked_group_zones = None   # 클릭한 그룹의 하위 구역 리스트
        self._current_emp_dict = {}       # 💡 현재 명단에 표시 중인 접속자 데이터
        self.debug_logs = []              # 💡 디버그 로그 (디버그 창 표시용)
        self.online_users = set()         # 💡 Client가 데이터를 보낸 사원 (PC ON 상태)

        # --- 그래프 클릭 매핑용 ---
        self._bar_objects = []
        self._bar_texts = []      # 💡 막대 옆 텍스트 객체 리스트 (호버/포커스 투명도용)
        self._bar_zone_map = []   # 각 bar에 대응하는 (zone_list, display_label)

        # 한글 폰트
        self.chart_font = None
        for font_name in ["Malgun Gothic", "맑은 고딕", "NanumGothic", "AppleGothic"]:
            if font_name in [f.name for f in fm.fontManager.ttflist]:
                self.chart_font = font_name
                break
        if self.chart_font:
            plt.rcParams['font.family'] = self.chart_font
        plt.rcParams['axes.unicode_minus'] = False

        self.settings = load_settings()
        self._update_job = None  # 💡 자동 갱신 타이머 ID

        # 💡 백그라운드 모니터링 / 트레이 연동 상태
        self.agent = None             # MonitoringAgent (main()에서 주입)
        self.tray_icon = None         # pystray Icon (main()에서 주입)
        self._quitting = False        # 완전 종료 진행 중 여부
        self._tray_notified = False   # 트레이 최초 안내 표시 여부

        self.setup_ui()
        self.root.protocol("WM_DELETE_WINDOW", self.on_close)  # 💡 종료 시 정리
        self.root.after(500, self.auto_connect_and_start)
        self.root.after(2000, self._poll_agent_status)  # 💡 내 PC 모니터링 상태 폴링

    def debug_log(self, msg):
        """콘솔 + 메모리에 로그 기록"""
        timestamp = datetime.now().strftime("%H:%M:%S")
        line = f"[{timestamp}] {msg}"
        print(line)
        self.debug_logs.append(line)
        if len(self.debug_logs) > 500:
            self.debug_logs.pop(0)

    def add_log_line(self, line):
        """💡 외부(모니터링 스레드)에서 이미 포맷된 로그 줄을 디버그 로그에 추가"""
        self.debug_logs.append(line)
        if len(self.debug_logs) > 500:
            self.debug_logs.pop(0)

    # ==========================================
    # 🔗 백그라운드 모니터링 / 트레이 연동
    # ==========================================
    def attach_agent(self, agent):
        self.agent = agent

    def attach_tray(self, tray_icon):
        self.tray_icon = tray_icon

    def _poll_agent_status(self):
        """내 PC 모니터링 상태를 상태바에 주기적으로 표시"""
        if self._quitting:
            return
        if self.agent is not None and hasattr(self, "lbl_monitor"):
            self.lbl_monitor.config(
                text=f"🖥 내 PC: {self.agent.summary}",
                fg=self.agent.summary_color,
            )
        self.root.after(3000, self._poll_agent_status)

    def show_window(self):
        """트레이에서 다시 창 열기"""
        try:
            self.root.deiconify()
            self.root.lift()
            self.root.focus_force()
        except Exception:
            pass

    def hide_to_tray(self):
        """창을 트레이로 숨김 (백그라운드 모니터링은 계속)"""
        try:
            self.root.withdraw()
        except Exception:
            pass
        if self.tray_icon is not None and not self._tray_notified:
            self._tray_notified = True
            try:
                self.tray_icon.notify(
                    "백그라운드에서 모니터링을 계속합니다.\n트레이 아이콘을 더블클릭하면 다시 열립니다.",
                    "한양이엔지 통합 관제",
                )
            except Exception:
                pass

    def real_quit(self):
        """프로그램 완전 종료 (모니터링 + 트레이 + GUI 정리)"""
        self._quitting = True

        # 자동 갱신 타이머 취소
        if self._update_job:
            try:
                self.root.after_cancel(self._update_job)
            except Exception:
                pass
            self._update_job = None

        # 백그라운드 모니터링 종료
        if self.agent is not None:
            try:
                self.agent.stop()
            except Exception:
                pass

        # 트레이 아이콘 종료
        if self.tray_icon is not None:
            try:
                self.tray_icon.stop()
            except Exception:
                pass

        # matplotlib 정리
        try:
            plt.close(self.fig)
        except Exception:
            pass

        try:
            self.root.destroy()
        except Exception:
            pass

    # ==========================================
    # UI 구성
    # ==========================================
    def setup_ui(self):
        # ── 상단 헤더 ──
        frame_header = tk.Frame(self.root, bg=self.colors["primary"], pady=8)
        frame_header.pack(fill=tk.X)

        tk.Label(frame_header, text="👤 한양이엔지 관제 시스템",
                 font=("맑은 고딕", 14, "bold"),
                 fg=self.colors["header_text"],
                 bg=self.colors["primary"]).pack(side=tk.LEFT, padx=15)

        btn_settings = tk.Button(frame_header, text="⚙ 설정", font=("맑은 고딕", 10),
                                  bg="#003C82", fg="white", relief=tk.FLAT, padx=10, pady=2,
                                  cursor="hand2", command=self.open_settings)
        btn_settings.pack(side=tk.RIGHT, padx=(5, 15))

        tk.Button(frame_header, text="🔍 디버그", font=("맑은 고딕", 10),
                  bg="#1B5E20", fg="white", relief=tk.FLAT, padx=10, pady=2,
                  cursor="hand2", command=self.open_debug_window).pack(side=tk.RIGHT, padx=(0, 5))

        # ── 상태바 ──
        frame_status = tk.Frame(self.root, bg=self.colors["secondary"], pady=6, padx=10)
        frame_status.pack(fill=tk.X)

        self.lbl_status = tk.Label(frame_status, text="⏳ 서버 연결 준비 중...",
                                    font=("맑은 고딕", 10, "bold"),
                                    bg=self.colors["secondary"], fg=self.colors["primary"])
        self.lbl_status.pack(side=tk.LEFT, padx=10)

        # 💡 내 PC 백그라운드 모니터링 상태 (구 Client.py 기능)
        self.lbl_monitor = tk.Label(frame_status, text="🖥 내 PC: 모니터링 준비 중…",
                                     font=("맑은 고딕", 9),
                                     bg=self.colors["secondary"], fg="#555555")
        self.lbl_monitor.pack(side=tk.LEFT, padx=(0, 10))

        self.lbl_time = tk.Label(frame_status, text="마지막 갱신: 대기 중",
                                  font=("맑은 고딕", 9),
                                  bg=self.colors["secondary"], fg="#757575")
        self.lbl_time.pack(side=tk.RIGHT, padx=(0, 10))

        # 💡 [수정] 갱신 주기 선택 + 수동 갱신 버튼
        saved_interval = self.settings.get("refresh_interval", "수동")
        self._refresh_interval_var = tk.StringVar(value=saved_interval)
        combo_refresh = ttk.Combobox(frame_status, textvariable=self._refresh_interval_var,
                                      values=["수동", "1분", "5분", "10분"],
                                      state="readonly", width=6, font=("맑은 고딕", 9))
        combo_refresh.pack(side=tk.RIGHT, padx=(0, 5))
        combo_refresh.bind("<<ComboboxSelected>>", self._on_refresh_interval_change)

        tk.Label(frame_status, text="갱신 주기:", font=("맑은 고딕", 9),
                 bg=self.colors["secondary"], fg="#555555").pack(side=tk.RIGHT, padx=(0, 3))

        tk.Button(frame_status, text="🔄 갱신", font=("맑은 고딕", 9, "bold"),
                  bg="#00479A", fg="white", relief=tk.FLAT, padx=10, pady=1,
                  cursor="hand2", command=self.manual_refresh).pack(side=tk.RIGHT, padx=(0, 5))

        tk.Button(frame_status, text="📥 현황 내려받기", font=("맑은 고딕", 9, "bold"),
                  bg="#2E7D32", fg="white", relief=tk.FLAT, padx=10, pady=1,
                  cursor="hand2", command=self.export_to_excel).pack(side=tk.RIGHT, padx=(0, 10))

        # ── 지역 버튼 ──
        frame_buttons = tk.Frame(self.root, bg=self.colors["background"], pady=12)
        frame_buttons.pack(fill=tk.X)

        tk.Label(frame_buttons, text="지역 선택 :",
                 font=("맑은 고딕", 11, "bold"),
                 bg=self.colors["background"], fg=self.colors["text"]).pack(side=tk.LEFT, padx=(20, 10))

        # 💡 [추가] 즐겨찾기 버튼 (가장 왼쪽)
        self.btn_favorites = tk.Button(
            frame_buttons, text="📌 즐겨찾기",
            font=("맑은 고딕", 11, "bold"),
            bg=self.colors["btn_normal"], fg=self.colors["btn_text_normal"],
            relief=tk.FLAT, padx=20, pady=6, cursor="hand2",
            command=self.show_favorites
        )
        self.btn_favorites.pack(side=tk.LEFT, padx=(0, 15))

        # 구분선
        tk.Frame(frame_buttons, bg="#CCCCCC", width=2).pack(side=tk.LEFT, fill=tk.Y, padx=(0, 15), pady=3)

        for btn_label in REGION_BUTTONS.keys():
            btn = tk.Button(
                frame_buttons, text=btn_label,
                font=("맑은 고딕", 11, "bold"),
                bg=self.colors["btn_normal"], fg=self.colors["btn_text_normal"],
                relief=tk.FLAT, padx=20, pady=6, cursor="hand2",
                command=lambda r=btn_label: self.on_region_click(r)
            )
            btn.pack(side=tk.LEFT, padx=5)
            self.region_buttons[btn_label] = btn

        # ── 메인 컨텐츠 (PanedWindow — 그래프/명단 사이 드래그로 크기 조절) ──
        self.frame_content = tk.PanedWindow(
            self.root, orient=tk.VERTICAL,
            bg="#AAAAAA", sashwidth=6, sashrelief=tk.RAISED,
            opaqueresize=True
        )
        self.frame_content.pack(fill=tk.BOTH, expand=True, padx=15, pady=(5, 10))

        # 그래프 카드
        self.frame_chart_area = tk.Frame(self.frame_content, bg=self.colors["card_bg"], relief=tk.RIDGE, bd=1)
        self.frame_content.add(self.frame_chart_area, minsize=200)

        self.lbl_chart_title = tk.Label(self.frame_chart_area, text="📊 라인별 접속 현황",
                                         font=("맑은 고딕", 12, "bold"),
                                         bg=self.colors["card_bg"], fg=self.colors["primary"])
        self.lbl_chart_title.pack(anchor=tk.W, padx=15, pady=(10, 0))

        # 범례 표시
        frame_legend = tk.Frame(self.frame_chart_area, bg=self.colors["card_bg"])
        frame_legend.pack(anchor=tk.E, padx=15)
        tk.Label(frame_legend, text="■", fg=self.colors["group_bar"], bg=self.colors["card_bg"],
                 font=("맑은 고딕", 10)).pack(side=tk.LEFT)
        tk.Label(frame_legend, text="그룹 합계  ", bg=self.colors["card_bg"],
                 font=("맑은 고딕", 9)).pack(side=tk.LEFT)
        tk.Label(frame_legend, text="■", fg=self.colors["sub_bar"], bg=self.colors["card_bg"],
                 font=("맑은 고딕", 10)).pack(side=tk.LEFT)
        tk.Label(frame_legend, text="개별 라인  ", bg=self.colors["card_bg"],
                 font=("맑은 고딕", 9)).pack(side=tk.LEFT)
        tk.Label(frame_legend, text="■", fg=self.colors["over_limit"], bg=self.colors["card_bg"],
                 font=("맑은 고딕", 10)).pack(side=tk.LEFT)
        tk.Label(frame_legend, text="제한 초과 ⚠  ", bg=self.colors["card_bg"],
                 font=("맑은 고딕", 9)).pack(side=tk.LEFT)
        tk.Label(frame_legend, text="(막대 클릭 → 접속자 확인)", bg=self.colors["card_bg"],
                 font=("맑은 고딕", 9), fg="#999999").pack(side=tk.LEFT)

        self.fig, self.ax = plt.subplots(figsize=(7, 4))
        self.fig.patch.set_facecolor("#FFFFFF")

        # 💡 캔버스 컨테이너 (줌 시 잔상 방지용)
        self.frame_canvas_container = tk.Frame(self.frame_chart_area, bg=self.colors["card_bg"])
        self.frame_canvas_container.pack(fill=tk.BOTH, expand=True, padx=10, pady=(5, 5))

        self.canvas_chart = FigureCanvasTkAgg(self.fig, master=self.frame_canvas_container)
        self.canvas_chart.get_tk_widget().pack(fill=tk.BOTH, expand=True)
        self.canvas_chart.mpl_connect("button_press_event", self.on_bar_click)
        self.canvas_chart.mpl_connect("scroll_event", self.on_scroll_zoom)
        self.canvas_chart.mpl_connect("motion_notify_event", self._on_bar_hover)

        # 💡 [추가] Ctrl + 마우스 휠 줌 기능
        self._zoom_level = 1.0          # 현재 줌 배율 (1.0 = 기본)
        self._zoom_min = 0.5            # 최소 축소
        self._zoom_max = 3.0            # 최대 확대
        self._base_figsize = (7, 4)     # 기본 그래프 크기
        self._hovered_bar_idx = -1      # 💡 현재 호버 중인 막대 인덱스
        self._focused_bar_idx = None    # 💡 클릭으로 포커스된 막대 인덱스 (None=전체 표시)

        # ── 접속자 명단 (PanedWindow의 하단 pane, 초기에는 추가하지 않음) ──
        self.frame_emp_area = tk.Frame(self.frame_content, bg=self.colors["card_bg"], relief=tk.RIDGE, bd=1)
        self._emp_panel_visible = False

        frame_emp_header = tk.Frame(self.frame_emp_area, bg=self.colors["card_bg"])
        frame_emp_header.pack(fill=tk.X, padx=10, pady=(8, 0))

        self.lbl_emp_title = tk.Label(frame_emp_header, text="👤 접속자 명단",
                                       font=("맑은 고딕", 11, "bold"),
                                       bg=self.colors["card_bg"], fg=self.colors["primary"])
        self.lbl_emp_title.pack(side=tk.LEFT)

        tk.Button(frame_emp_header, text="✕ 닫기", font=("맑은 고딕", 9),
                  bg="#EEEEEE", fg="#555555", relief=tk.FLAT, padx=8, pady=2,
                  cursor="hand2", command=self.hide_emp_list).pack(side=tk.RIGHT)

        # 💡 [수정] 좌우 분할: 왼쪽 = 팀/공종별 인원, 오른쪽 = 사원명 목록
        frame_emp_body = tk.PanedWindow(self.frame_emp_area, orient=tk.HORIZONTAL,
                                         bg="#CCCCCC", sashwidth=4, sashrelief=tk.RAISED)
        frame_emp_body.pack(fill=tk.BOTH, expand=True, padx=10, pady=(5, 8))

        # --- 왼쪽: 팀/공종 트리 ---
        frame_left_emp = tk.Frame(frame_emp_body, bg=self.colors["card_bg"])
        frame_emp_body.add(frame_left_emp, minsize=200)

        tk.Label(frame_left_emp, text="📂 팀 / 공종별 인원",
                 font=("맑은 고딕", 10, "bold"),
                 bg=self.colors["card_bg"], fg=self.colors["primary"]).pack(anchor=tk.W, padx=5, pady=(3, 3))

        scroll_left = ttk.Scrollbar(frame_left_emp, orient=tk.VERTICAL)
        self.tree_team = ttk.Treeview(frame_left_emp, columns=("Count",),
                                       show="tree headings", height=8,
                                       yscrollcommand=scroll_left.set)
        scroll_left.config(command=self.tree_team.yview)
        scroll_left.pack(side=tk.RIGHT, fill=tk.Y)

        self.tree_team.heading("#0", text="팀 / 공종", anchor=tk.W)
        self.tree_team.heading("Count", text="인원", anchor="center")
        self.tree_team.column("#0", width=180, anchor=tk.W)
        self.tree_team.column("Count", width=60, anchor="center")
        self.tree_team.pack(fill=tk.BOTH, expand=True)
        self.tree_team.bind("<<TreeviewSelect>>", self._on_team_tree_select)

        # --- 오른쪽: 사원명 목록 ---
        frame_right_emp = tk.Frame(frame_emp_body, bg=self.colors["card_bg"])
        frame_emp_body.add(frame_right_emp, minsize=200)

        self.lbl_name_title = tk.Label(frame_right_emp, text="👤 사원 목록 (공종 선택)",
                                        font=("맑은 고딕", 10, "bold"),
                                        bg=self.colors["card_bg"], fg=self.colors["primary"])
        self.lbl_name_title.pack(anchor=tk.W, padx=5, pady=(3, 3))

        scroll_right = ttk.Scrollbar(frame_right_emp, orient=tk.VERTICAL)
        self.tree_names = ttk.Treeview(frame_right_emp, columns=("Name",),
                                        show="headings", height=8,
                                        yscrollcommand=scroll_right.set)
        scroll_right.config(command=self.tree_names.yview)
        scroll_right.pack(side=tk.RIGHT, fill=tk.Y)

        self.tree_names.heading("Name", text="사원명", anchor="center")
        self.tree_names.column("Name", width=250, anchor="center")
        self.tree_names.pack(fill=tk.BOTH, expand=True)

        style = ttk.Style()
        style.configure("Treeview", rowheight=30, font=("맑은 고딕", 10))
        style.configure("Treeview.Heading", font=("맑은 고딕", 10, "bold"),
                         foreground=self.colors["primary"])

        self.draw_empty_chart("서버 연결 대기 중...")

    # ==========================================
    # 📌 즐겨찾기 라인 (시작 화면)
    # ==========================================
    def show_favorites(self):
        """아무 지역도 선택하지 않았을 때 즐겨찾기 라인만 표시"""
        self.selected_region = None
        self.clicked_zone = None
        self.clicked_group_zones = None
        self._reset_zoom()
        self.hide_emp_list()

        # 모든 지역 버튼 비활성 상태로, 즐겨찾기 버튼만 활성
        for label, btn in self.region_buttons.items():
            btn.config(bg=self.colors["btn_normal"], fg=self.colors["btn_text_normal"])
        self.btn_favorites.config(bg=self.colors["btn_active"], fg=self.colors["btn_text_active"])

        fav_zones = self.settings.get("favorite_zones", [])
        if not fav_zones:
            self.draw_empty_chart("설정에서 즐겨찾기 라인을 추가하세요")
            return

        self._draw_favorites_chart(fav_zones)

    def _draw_favorites_chart(self, fav_zones):
        """즐겨찾기 라인들만 모아서 평면 막대그래프"""
        self.ax.clear()
        self._bar_objects = []
        self._bar_texts = []
        self._bar_zone_map = []

        display_labels = [get_display_name(z) for z in fav_zones]
        counts = []
        for zone in fav_zones:
            count = len(self._count_zone_users(zone))
            counts.append(count)

        bar_colors = [self._get_bar_color(z, c) for z, c in zip(fav_zones, counts)]
        bars = self.ax.barh(range(len(fav_zones)), counts, color=bar_colors, height=0.6, edgecolor="white")
        self._bar_objects = list(bars)

        for zone in fav_zones:
            self._bar_zone_map.append(([zone], get_display_name(zone)))

        max_val = max(counts) if counts and max(counts) > 0 else 1
        for bar, count, zone in zip(bars, counts, fav_zones):
            x_pos = bar.get_width() + max_val * 0.02
            over = self._is_over_limit(zone, count)
            limit_val = self.server_zone_limits.get(zone)
            label = self._format_count_label(count, limit_val, over)
            text_color = self.colors["over_limit"] if over else (self.colors["primary"] if count > 0 else "#999999")
            txt = self.ax.text(x_pos, bar.get_y() + bar.get_height() / 2,
                         label, va='center', ha='left',
                         fontsize=10, fontweight='bold', color=text_color)
            self._bar_texts.append(txt)

        self.ax.set_yticks(range(len(fav_zones)))
        self.ax.set_yticklabels(display_labels)
        for label_obj in self.ax.get_yticklabels():
            label_obj.set_fontsize(11)
            label_obj.set_color(self.colors["text"])

        max_val = max(counts) if counts and max(counts) > 0 else 5
        self.ax.set_xlim(0, max_val * 1.3)
        self.ax.invert_yaxis()
        self.ax.set_xlabel("접속 인원 (명)", fontsize=9, color="#666666")
        self.ax.tick_params(axis='x', labelsize=9)
        self.ax.spines['top'].set_visible(False)
        self.ax.spines['right'].set_visible(False)
        self.ax.set_facecolor("#FAFAFA")

        total = sum(counts)
        self.lbl_chart_title.config(text=f"📊 라인별 접속 현황 — 총 {total}명 접속 중")

        # 💡 캔버스 재생성으로 잔상 완전 제거
        self._rebuild_canvas()
    def open_settings(self):
        dlg = tk.Toplevel(self.root)
        dlg.title("⚙ 설정")
        dlg.geometry("500x520")
        dlg.resizable(False, False)
        dlg.configure(bg=self.colors["card_bg"])
        dlg.transient(self.root)
        dlg.grab_set()

        dlg.update_idletasks()
        x = self.root.winfo_x() + (self.root.winfo_width() - 500) // 2
        y = self.root.winfo_y() + (self.root.winfo_height() - 520) // 2
        dlg.geometry(f"+{x}+{y}")

        tk.Label(dlg, text="⚙ 뷰어 설정", font=("맑은 고딕", 13, "bold"),
                 bg=self.colors["card_bg"], fg=self.colors["primary"]).pack(pady=(12, 5))

        # ── 탭 구성 ──
        notebook = ttk.Notebook(dlg)
        notebook.pack(fill=tk.BOTH, expand=True, padx=15, pady=5)

        # ======== 탭 1: 즐겨찾기 + 기본 지역 ========
        tab_general = tk.Frame(notebook, bg=self.colors["card_bg"])
        notebook.add(tab_general, text="  📌 즐겨찾기 / 기본 지역  ")

        tk.Label(tab_general, text="📌 시작 화면 즐겨찾기 라인",
                 font=("맑은 고딕", 10, "bold"),
                 bg=self.colors["card_bg"], fg=self.colors["primary"]).pack(anchor=tk.W, padx=15, pady=(10, 2))
        tk.Label(tab_general, text="아무 지역도 선택하지 않았을 때 보여줄 라인입니다. (순서 변경 가능)",
                 font=("맑은 고딕", 9),
                 bg=self.colors["card_bg"], fg="#888888").pack(anchor=tk.W, padx=15)

        # 💡 [수정] 즐겨찾기 — 순서 변경 가능한 리스트
        all_zones = self._get_all_zones()
        all_zone_map = {z[0]: f"{z[1]}  ({z[2]})" for z in all_zones}  # 내부명 → 표시명
        current_favs = list(self.settings.get("favorite_zones", []))

        frame_fav_body = tk.Frame(tab_general, bg=self.colors["card_bg"])
        frame_fav_body.pack(fill=tk.X, padx=15, pady=(5, 5))

        # 추가 영역 (상단)
        frame_fav_add = tk.Frame(frame_fav_body, bg=self.colors["card_bg"])
        frame_fav_add.pack(fill=tk.X, pady=(0, 5))

        available_items = [f"{z[1]}  ({z[2]})" for z in all_zones]
        combo_add_var = tk.StringVar()
        combo_add = ttk.Combobox(frame_fav_add, textvariable=combo_add_var,
                                  values=available_items, state="readonly",
                                  width=30, font=("맑은 고딕", 9))
        combo_add.pack(side=tk.LEFT, padx=(0, 5))

        # 리스트 + 버튼 영역
        frame_fav_list = tk.Frame(frame_fav_body, bg=self.colors["card_bg"])
        frame_fav_list.pack(fill=tk.X)

        fav_listbox = tk.Listbox(frame_fav_list, height=6, font=("맑은 고딕", 9),
                                  selectmode=tk.SINGLE, activestyle="none")
        fav_listbox.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)

        # 현재 즐겨찾기 채우기
        fav_internal_list = []  # 내부명 리스트 (순서 유지)
        for zone_key in current_favs:
            display = all_zone_map.get(zone_key, get_display_name(zone_key))
            fav_listbox.insert(tk.END, display)
            fav_internal_list.append(zone_key)

        # 오른쪽 버튼들
        frame_fav_btns = tk.Frame(frame_fav_list, bg=self.colors["card_bg"])
        frame_fav_btns.pack(side=tk.LEFT, padx=(5, 0))

        def fav_add():
            sel_text = combo_add_var.get()
            if not sel_text: return
            # 표시명 → 내부명 역변환
            for z_int, z_disp in all_zone_map.items():
                if z_disp == sel_text:
                    if z_int not in fav_internal_list:
                        fav_internal_list.append(z_int)
                        fav_listbox.insert(tk.END, sel_text)
                    break

        def fav_remove():
            sel = fav_listbox.curselection()
            if not sel: return
            idx = sel[0]
            fav_listbox.delete(idx)
            fav_internal_list.pop(idx)

        def fav_up():
            sel = fav_listbox.curselection()
            if not sel or sel[0] == 0: return
            idx = sel[0]
            # 리스트 swap
            fav_internal_list[idx], fav_internal_list[idx-1] = fav_internal_list[idx-1], fav_internal_list[idx]
            text = fav_listbox.get(idx)
            fav_listbox.delete(idx)
            fav_listbox.insert(idx-1, text)
            fav_listbox.selection_set(idx-1)

        def fav_down():
            sel = fav_listbox.curselection()
            if not sel or sel[0] >= fav_listbox.size()-1: return
            idx = sel[0]
            fav_internal_list[idx], fav_internal_list[idx+1] = fav_internal_list[idx+1], fav_internal_list[idx]
            text = fav_listbox.get(idx)
            fav_listbox.delete(idx)
            fav_listbox.insert(idx+1, text)
            fav_listbox.selection_set(idx+1)

        tk.Button(frame_fav_btns, text="추가 ←", font=("맑은 고딕", 9),
                  width=7, command=fav_add).pack(pady=2)
        tk.Button(frame_fav_btns, text="제거 ✕", font=("맑은 고딕", 9),
                  width=7, command=fav_remove).pack(pady=2)
        tk.Button(frame_fav_btns, text="▲ 위로", font=("맑은 고딕", 9),
                  width=7, command=fav_up).pack(pady=2)
        tk.Button(frame_fav_btns, text="▼ 아래로", font=("맑은 고딕", 9),
                  width=7, command=fav_down).pack(pady=2)

        ttk.Separator(tab_general, orient=tk.HORIZONTAL).pack(fill=tk.X, padx=15, pady=8)

        frame_opt = tk.Frame(tab_general, bg=self.colors["card_bg"])
        frame_opt.pack(fill=tk.X, padx=15, pady=2)

        tk.Label(frame_opt, text="🏠 기본 시작 지역 :",
                 font=("맑은 고딕", 10, "bold"),
                 bg=self.colors["card_bg"], fg=self.colors["primary"]).pack(side=tk.LEFT)

        combo_var = tk.StringVar(value=self._get_button_label(self.settings.get("default_region", "평택")))
        combo = ttk.Combobox(frame_opt, textvariable=combo_var,
                              values=["즐겨찾기 (버튼 없음)"] + list(REGION_BUTTONS.keys()),
                              state="readonly", width=18, font=("맑은 고딕", 10))
        if self.settings.get("default_region") == "__favorites__":
            combo_var.set("즐겨찾기 (버튼 없음)")
        combo.pack(side=tk.LEFT, padx=10)

        # ======== 탭 2: 접속 제한 현황 (읽기 전용 — 서버에서 관리) ========
        tab_limits = tk.Frame(notebook, bg=self.colors["card_bg"])
        notebook.add(tab_limits, text="  🚨 접속 제한 현황  ")

        tk.Label(tab_limits, text="🚨 라인별 접속 제한 인원 (서버에서 관리)",
                 font=("맑은 고딕", 10, "bold"),
                 bg=self.colors["card_bg"], fg=self.colors["primary"]).pack(anchor=tk.W, padx=15, pady=(10, 2))
        tk.Label(tab_limits, text="이 설정은 최고관리자 전용 앱에서만 변경할 수 있습니다.",
                 font=("맑은 고딕", 9),
                 bg=self.colors["card_bg"], fg="#888888").pack(anchor=tk.W, padx=15)

        # 접속 제한 스크롤 영역
        frame_limits_outer = tk.Frame(tab_limits, bg=self.colors["card_bg"], relief=tk.SUNKEN, bd=1)
        frame_limits_outer.pack(fill=tk.BOTH, expand=True, padx=15, pady=(5, 10))

        canvas_limits = tk.Canvas(frame_limits_outer, bg=self.colors["card_bg"], highlightthickness=0)
        scrollbar_limits = ttk.Scrollbar(frame_limits_outer, orient=tk.VERTICAL, command=canvas_limits.yview)
        frame_limits_inner = tk.Frame(canvas_limits, bg=self.colors["card_bg"])

        frame_limits_inner.bind("<Configure>",
                                lambda e: canvas_limits.configure(scrollregion=canvas_limits.bbox("all")))
        canvas_limits.create_window((0, 0), window=frame_limits_inner, anchor="nw")
        canvas_limits.configure(yscrollcommand=scrollbar_limits.set)
        canvas_limits.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar_limits.pack(side=tk.RIGHT, fill=tk.Y)

        # 헤더
        frame_limit_hdr = tk.Frame(frame_limits_inner, bg="#E8EAF6")
        frame_limit_hdr.pack(fill=tk.X, padx=2, pady=(2, 0))
        tk.Label(frame_limit_hdr, text="구역", width=12, font=("맑은 고딕", 9, "bold"),
                 bg="#E8EAF6", fg=self.colors["primary"]).pack(side=tk.LEFT, padx=5)
        tk.Label(frame_limit_hdr, text="지역", width=14, font=("맑은 고딕", 9, "bold"),
                 bg="#E8EAF6", fg=self.colors["primary"]).pack(side=tk.LEFT, padx=5)
        tk.Label(frame_limit_hdr, text="제한 인원", width=10, font=("맑은 고딕", 9, "bold"),
                 bg="#E8EAF6", fg=self.colors["primary"]).pack(side=tk.LEFT, padx=5)

        current_limits = self.server_zone_limits

        for zone_internal, display_name, region_label in all_zones:
            frame_row = tk.Frame(frame_limits_inner, bg=self.colors["card_bg"])
            frame_row.pack(fill=tk.X, padx=2, pady=1)

            tk.Label(frame_row, text=display_name, width=12, anchor=tk.W,
                     font=("맑은 고딕", 9), bg=self.colors["card_bg"]).pack(side=tk.LEFT, padx=5)
            tk.Label(frame_row, text=region_label, width=14, anchor=tk.W,
                     font=("맑은 고딕", 9), bg=self.colors["card_bg"], fg="#888888").pack(side=tk.LEFT, padx=5)

            # 💡 읽기 전용 — 서버값만 표시
            limit_val = current_limits.get(zone_internal)
            display_val = f"{limit_val}명" if limit_val else "제한 없음"
            color = self.colors["status_offline"] if limit_val else "#999999"
            tk.Label(frame_row, text=display_val, width=10,
                     font=("맑은 고딕", 9, "bold"), fg=color,
                     bg=self.colors["card_bg"]).pack(side=tk.LEFT, padx=5)

        # ======== 저장 버튼 (공통) ========
        def on_save():
            # 즐겨찾기 저장
            new_favs = list(fav_internal_list)  # 순서 유지된 즐겨찾기 리스트
            self.settings["favorite_zones"] = new_favs

            # 기본 지역 저장
            selected_label = combo_var.get()
            if selected_label == "즐겨찾기 (버튼 없음)":
                self.settings["default_region"] = "__favorites__"
            else:
                internal_key = REGION_BUTTONS.get(selected_label, "평택")
                self.settings["default_region"] = internal_key

            save_settings(self.settings)
            dlg.destroy()

            # 현재 화면 갱신
            if not self.selected_region:
                self.show_favorites()
            else:
                self.refresh_chart()

        tk.Button(dlg, text="💾 저장", font=("맑은 고딕", 11, "bold"),
                  bg=self.colors["primary"], fg="white",
                  relief=tk.FLAT, padx=25, pady=5,
                  cursor="hand2", command=on_save).pack(pady=10)

    def _get_all_zones(self):
        """모든 지역의 구역 목록 반환: [(내부명, 표시명, 지역라벨), ...]"""
        result = []
        for btn_label, internal_region in REGION_BUTTONS.items():
            config = REGION_CONFIG.get(internal_region, {})
            if config.get("grouped"):
                for group_name, sub_zones in config.get("groups", {}).items():
                    for zone in sub_zones:
                        result.append((zone, get_display_name(zone), f"{btn_label} > {group_name}"))
            else:
                for zone in config.get("zones", []):
                    result.append((zone, get_display_name(zone), btn_label))
        return result

    def _get_button_label(self, internal_key):
        for label, key in REGION_BUTTONS.items():
            if key == internal_key:
                return label
        return "평택"

    # ==========================================
    # 🖱 지역 버튼 클릭
    # ==========================================
    def on_region_click(self, btn_label):
        self.selected_region = btn_label
        self.clicked_zone = None
        self.clicked_group_zones = None
        self._reset_zoom()
        self.hide_emp_list()

        for label, btn in self.region_buttons.items():
            if label == btn_label:
                btn.config(bg=self.colors["btn_active"], fg=self.colors["btn_text_active"])
            else:
                btn.config(bg=self.colors["btn_normal"], fg=self.colors["btn_text_normal"])
        self.btn_favorites.config(bg=self.colors["btn_normal"], fg=self.colors["btn_text_normal"])

        self.refresh_chart()

    # ==========================================
    # 📊 가로 막대 그래프
    # ==========================================
    def _count_zone_users(self, zone_internal):
        """특정 구역의 접속 사원 set 반환
        💡 "구역:프로세스명" 형태면 해당 프로세스만 필터링
           "구역" 형태면 모든 프로그램 합산
        """
        # 프로세스 필터 분리
        if ":" in zone_internal:
            actual_zone, proc_filter = zone_internal.split(":", 1)
        else:
            actual_zone, proc_filter = zone_internal, None

        users = set()
        for (emp_name, proc_name), user_zones in self.active_users.items():
            # 프로세스 필터가 있으면 해당 프로세스만
            if proc_filter and proc_name.lower() != proc_filter.lower():
                continue
            for uz in user_zones:
                if uz == actual_zone:
                    users.add(emp_name)
                    break
        return users

    def _is_over_limit(self, zone_internal, count):
        """💡 해당 구역이 제한 인원을 초과했는지 확인"""
        limits = self.server_zone_limits
        if zone_internal in limits:
            return count > limits[zone_internal]
        return False

    def _format_count_label(self, count, limit_val, is_over):
        """💡 막대 옆 텍스트 포맷
        제한 있음: "3/5명" 또는 "6/5명 ⚠" (초과 시)
        제한 없음: "3명"
        """
        if limit_val is not None:
            label = f'{count}/{limit_val}명'
            if is_over:
                label += ' ⚠'
            return label
        return f'{count}명'

    def _get_bar_color(self, zone_internal, count, is_group=False):
        """💡 막대 색상 결정: 제한 초과 → 빨간색, 일반 → 기본색"""
        if count == 0:
            return self.colors["empty_bar"]
        if self._is_over_limit(zone_internal, count):
            return self.colors["over_limit_group"] if is_group else self.colors["over_limit"]
        return self.colors["group_bar"] if is_group else self.colors["sub_bar"]

    def _get_group_bar_color(self, sub_zones, group_total):
        """💡 그룹 막대 색상: 하위 라인 중 하나라도 초과하면 빨간색"""
        limits = self.server_zone_limits
        for zone in sub_zones:
            if zone in limits:
                zone_count = len(self._count_zone_users(zone))
                if zone_count > limits[zone]:
                    return self.colors["over_limit_group"]
        if group_total == 0:
            return self.colors["empty_bar"]
        return self.colors["group_bar"]

    def refresh_chart(self):
        if not self.selected_region:
            # 💡 [기능1] 즐겨찾기 모드
            fav_zones = self.settings.get("favorite_zones", [])
            if fav_zones:
                self._draw_favorites_chart(fav_zones)
            else:
                self.draw_empty_chart("설정에서 즐겨찾기 라인을 추가하세요")
            return

        internal_region = REGION_BUTTONS.get(self.selected_region, "")
        config = REGION_CONFIG.get(internal_region, {})

        if config.get("grouped"):
            self._draw_grouped_chart(config)
        else:
            self._draw_flat_chart(config)
            # 💡 [기능2] 라인이 적은 지역 — 처음 진입 시에만 전체 명단 표시
            # 갱신 시에는 현재 선택(특정 막대 또는 전체)을 유지
            if not self._emp_panel_visible:
                all_zones = config.get("zones", [])
                if all_zones:
                    self.show_emp_list_for_zones(all_zones, self.selected_region)

    def _draw_grouped_chart(self, config):
        """평택용 - 그룹(P1,P2...) + 하위 라인 계층형 막대그래프"""
        groups = config.get("groups", {})
        if not groups:
            self.draw_empty_chart(f"'{self.selected_region}' 등록된 라인이 없습니다")
            return

        self.ax.clear()
        self._bar_objects = []
        self._bar_texts = []
        self._bar_zone_map = []

        y_labels = []
        y_values = []
        y_colors = []
        y_heights = []
        y_is_group = []
        y_positions = []

        current_y = 0
        group_gap = 0.4   # 그룹 사이 간격

        for group_name, sub_zones in groups.items():
            # 그룹 합계 계산
            group_users = set()
            for zone in sub_zones:
                group_users |= self._count_zone_users(zone)
            group_total = len(group_users)

            # 그룹 행
            y_positions.append(current_y)
            y_labels.append(f"■ {group_name}")
            y_values.append(group_total)
            y_colors.append(self._get_group_bar_color(sub_zones, group_total))
            y_heights.append(0.7)
            y_is_group.append(True)
            self._bar_zone_map.append((sub_zones, group_name))
            current_y += 1

            # 하위 라인 행
            for zone in sub_zones:
                zone_users = self._count_zone_users(zone)
                count = len(zone_users)

                y_positions.append(current_y)
                y_labels.append(f"    {get_display_name(zone)}")
                y_values.append(count)
                y_colors.append(self._get_bar_color(zone, count))
                y_heights.append(0.5)
                y_is_group.append(False)
                self._bar_zone_map.append(([zone], get_display_name(zone)))
                current_y += 1

            current_y += group_gap  # 그룹 사이 간격

        # 그래프 그리기 (위→아래 순서)
        bars = self.ax.barh(y_positions, y_values, color=y_colors, height=y_heights, edgecolor="white")
        self._bar_objects = list(bars)

        # 인원수 텍스트
        max_val = max(y_values) if y_values and max(y_values) > 0 else 1
        for i, (bar, val, is_grp) in enumerate(zip(bars, y_values, y_is_group)):
            x_pos = bar.get_width() + max_val * 0.02
            zone_list, disp = self._bar_zone_map[i]

            # 💡 제한 초과 여부 확인
            has_over = False
            if is_grp:
                # 그룹: 하위 라인 중 하나라도 초과하면 경고
                limits = self.server_zone_limits
                for z in zone_list:
                    if z in limits and len(self._count_zone_users(z)) > limits[z]:
                        has_over = True
                        break
                label = f'합계 {val}명' + (' ⚠' if has_over else '')
            else:
                zone = zone_list[0]
                has_over = self._is_over_limit(zone, val)
                limit_val = self.server_zone_limits.get(zone)
                label = self._format_count_label(val, limit_val, has_over)

            text_color = self.colors["over_limit"] if has_over else (self.colors["group_bar"] if val > 0 else "#999999")
            txt = self.ax.text(x_pos, bar.get_y() + bar.get_height() / 2,
                         label, va='center', ha='left',
                         fontsize=9 if is_grp else 8,
                         fontweight='bold' if is_grp else 'normal',
                         color=text_color)
            self._bar_texts.append(txt)

        self.ax.set_yticks(y_positions)
        self.ax.set_yticklabels(y_labels)

        # 그룹 라벨 볼드 처리
        for i, (label_obj, is_grp) in enumerate(zip(self.ax.get_yticklabels(), y_is_group)):
            if is_grp:
                label_obj.set_fontweight('bold')
                label_obj.set_fontsize(11)
                label_obj.set_color(self.colors["group_bar"])
            else:
                label_obj.set_fontsize(10)
                label_obj.set_color(self.colors["text"])

        self._finalize_chart(y_values)

    def _draw_flat_chart(self, config):
        """기흥•화성, 천안, 해외용 - 평면 막대그래프"""
        zones = config.get("zones", [])
        if not zones:
            self.draw_empty_chart(f"'{self.selected_region}' 등록된 라인이 없습니다")
            return

        self.ax.clear()
        self._bar_objects = []
        self._bar_texts = []
        self._bar_zone_map = []

        display_labels = [get_display_name(z) for z in zones]
        counts = []
        for zone in zones:
            count = len(self._count_zone_users(zone))
            counts.append(count)

        bar_colors = [self._get_bar_color(z, c) for z, c in zip(zones, counts)]
        bars = self.ax.barh(range(len(zones)), counts, color=bar_colors, height=0.6, edgecolor="white")
        self._bar_objects = list(bars)

        for zone in zones:
            self._bar_zone_map.append(([zone], get_display_name(zone)))

        max_val = max(counts) if counts and max(counts) > 0 else 1
        for bar, count, zone in zip(bars, counts, zones):
            x_pos = bar.get_width() + max_val * 0.02
            over = self._is_over_limit(zone, count)
            limit_val = self.server_zone_limits.get(zone)
            label = self._format_count_label(count, limit_val, over)
            text_color = self.colors["over_limit"] if over else (self.colors["primary"] if count > 0 else "#999999")
            txt = self.ax.text(x_pos, bar.get_y() + bar.get_height() / 2,
                         label, va='center', ha='left',
                         fontsize=10, fontweight='bold', color=text_color)
            self._bar_texts.append(txt)

        self.ax.set_yticks(range(len(zones)))
        self.ax.set_yticklabels(display_labels)
        for label_obj in self.ax.get_yticklabels():
            label_obj.set_fontsize(11)
            label_obj.set_color(self.colors["text"])

        self._finalize_chart(counts)

    def _finalize_chart(self, values):
        """그래프 공통 마무리"""
        max_val = max(values) if values and max(values) > 0 else 5
        self.ax.set_xlim(0, max_val * 1.3)
        self.ax.invert_yaxis()
        self.ax.set_xlabel("접속 인원 (명)", fontsize=9, color="#666666")
        self.ax.tick_params(axis='x', labelsize=9)
        self.ax.spines['top'].set_visible(False)
        self.ax.spines['right'].set_visible(False)
        self.ax.set_facecolor("#FAFAFA")

        total_users = set()
        internal_region = REGION_BUTTONS.get(self.selected_region, "")
        cfg = REGION_CONFIG.get(internal_region, {})
        all_zone_keys = []
        if cfg.get("grouped"):
            for sub_zones in cfg.get("groups", {}).values():
                all_zone_keys.extend(sub_zones)
        else:
            all_zone_keys = cfg.get("zones", [])

        # 💡 "구역:프로세스" 형태에서 실제 구역명만 추출
        actual_zones = set()
        for zk in all_zone_keys:
            actual_zones.add(zk.split(":")[0])

        for (emp_name, proc_name), zones in self.active_users.items():
            for uz in zones:
                if uz in actual_zones:
                    total_users.add(emp_name)
                    break

        self.lbl_chart_title.config(
            text=f"📊 [{self.selected_region}] 라인별 접속 현황 — 총 {len(total_users)}명 접속 중"
        )

        # 💡 캔버스 재생성으로 잔상 완전 제거
        self._rebuild_canvas()

    def draw_empty_chart(self, message):
        self.ax.clear()
        self.ax.text(0.5, 0.5, message, transform=self.ax.transAxes,
                     fontsize=13, ha='center', va='center', color="#999999")
        for spine in self.ax.spines.values():
            spine.set_visible(False)
        self.ax.set_xticks([])
        self.ax.set_yticks([])
        self.ax.set_facecolor("#FAFAFA")
        self.lbl_chart_title.config(text="📊 라인별 접속 현황")
        self._bar_objects = []
        self._bar_texts = []
        self._bar_zone_map = []
        # 💡 캔버스 재생성으로 잔상 완전 제거
        self._rebuild_canvas()

    # ==========================================
    # 🔍 Ctrl + 마우스 휠 → 그래프 줌
    # ==========================================
    def on_scroll_zoom(self, event):
        """Ctrl 키를 누른 상태에서 마우스 휠로 그래프 확대/축소"""
        if event.inaxes != self.ax:
            return
        if event.key != 'control':
            return

        if event.button == 'up':
            self._zoom_level = min(self._zoom_level + 0.15, self._zoom_max)
        elif event.button == 'down':
            self._zoom_level = max(self._zoom_level - 0.15, self._zoom_min)

        self._rebuild_canvas()

    def _rebuild_canvas(self):
        """💡 캔버스를 파괴하고 새로 생성 — 잔상 완전 제거
        줌 변경, 자동 갱신, 지역 전환 등 모든 차트 갱신에서 공통 사용
        """
        new_w = self._base_figsize[0]
        new_h = self._base_figsize[1] * self._zoom_level
        self.fig.set_size_inches(new_w, new_h)
        self.fig.subplots_adjust(left=0.18, right=0.92, top=0.95, bottom=0.12)

        # 기존 캔버스 완전 제거
        self.canvas_chart.get_tk_widget().destroy()

        # 새 캔버스 생성
        self.canvas_chart = FigureCanvasTkAgg(self.fig, master=self.frame_canvas_container)
        self.canvas_chart.get_tk_widget().pack(fill=tk.BOTH, expand=True)
        self.canvas_chart.mpl_connect("button_press_event", self.on_bar_click)
        self.canvas_chart.mpl_connect("scroll_event", self.on_scroll_zoom)
        self.canvas_chart.mpl_connect("motion_notify_event", self._on_bar_hover)
        self.canvas_chart.draw()

    def _reset_zoom(self):
        """줌 배율을 기본값으로 초기화"""
        self._zoom_level = 1.0
        self._rebuild_canvas()

    # ==========================================
    # 🖱 막대 클릭 → 접속자 명단
    # ==========================================
    def _on_bar_hover(self, event):
        """💡 마우스 호버 시 막대 강조 효과"""
        if not self._bar_objects or event.inaxes != self.ax:
            if self._hovered_bar_idx >= 0:
                self._hovered_bar_idx = -1
                self._update_bar_highlight()
            return

        new_hover = -1
        for i, bar in enumerate(self._bar_objects):
            if bar.contains(event)[0]:
                new_hover = i
                break

        if new_hover != self._hovered_bar_idx:
            self._hovered_bar_idx = new_hover
            self._update_bar_highlight()

    def _update_bar_highlight(self):
        """호버/포커스 상태에 따라 막대 + 텍스트 투명도 업데이트"""
        if not self._bar_objects:
            return

        for i, bar in enumerate(self._bar_objects):
            # 텍스트 객체 (있으면)
            txt = self._bar_texts[i] if i < len(self._bar_texts) else None

            if self._focused_bar_idx is not None:
                # 포커스 모드: 선택된 바만 보이고 나머지는 거의 투명
                if i == self._focused_bar_idx:
                    alpha = 1.0
                else:
                    alpha = 0.08
            elif self._hovered_bar_idx >= 0:
                # 호버 모드: 호버된 바 강조, 나머지 살짝 투명
                if i == self._hovered_bar_idx:
                    alpha = 1.0
                else:
                    alpha = 0.5
            else:
                # 기본: 모두 불투명
                alpha = 1.0

            bar.set_alpha(alpha)
            if txt:
                txt.set_alpha(alpha)

        self.canvas_chart.draw_idle()

    def on_bar_click(self, event):
        if not self._bar_objects or event.inaxes != self.ax:
            return

        for i, bar in enumerate(self._bar_objects):
            if bar.contains(event)[0]:
                zone_list, display_label = self._bar_zone_map[i]
                # 💡 [기능4] 클릭한 바만 강조, 나머지 숨김
                self._focused_bar_idx = i
                self._update_bar_highlight()
                self.show_emp_list_for_zones(zone_list, display_label)
                return

    def show_emp_list_for_zones(self, zone_list, display_label):
        self.clicked_zone = zone_list[0] if len(zone_list) == 1 else None
        self.clicked_group_zones = zone_list

        # 💡 접속자 명단 표시 시 창 높이 확장
        current_w = self.root.winfo_width()
        current_h = self.root.winfo_height()
        min_h = 900
        if current_h < min_h:
            self.root.geometry(f"{current_w}x{min_h}")
        self.root.minsize(950, 900)  # 💡 명단 표시 중일 때 최소 높이 확장

        # PanedWindow에 명단 패널 추가
        if not self._emp_panel_visible:
            self.frame_content.add(self.frame_emp_area, minsize=200)
            self._emp_panel_visible = True

        # 💡 접속자 수집 (사원명 기준 중복 제거)
        self._current_emp_dict = {}  # 사원명 → {"progs": set, "zones": set, "team": str, "gongjong": str}
        for (emp_name, proc_name), user_zones in self.active_users.items():
            for zone_key in zone_list:
                if ":" in zone_key:
                    actual_zone, proc_filter = zone_key.split(":", 1)
                    if proc_name.lower() != proc_filter.lower():
                        continue
                else:
                    actual_zone = zone_key

                for uz in user_zones:
                    if uz == actual_zone:
                        proc_label = PROGRAMS_MAP.get(proc_name.lower(), proc_name)
                        if emp_name not in self._current_emp_dict:
                            info = self.emp_info.get(emp_name, {"team": "미지정", "gongjong": "미지정"})
                            self._current_emp_dict[emp_name] = {
                                "progs": set(), "zones": set(),
                                "team": info["team"], "gongjong": info["gongjong"]
                            }
                        self._current_emp_dict[emp_name]["progs"].add(proc_label)
                        self._current_emp_dict[emp_name]["zones"].add(get_display_name(zone_key))
                        break

        # 💡 왼쪽 트리: 팀 → 공종별 인원수
        for item in self.tree_team.get_children():
            self.tree_team.delete(item)

        team_gongjong = {}  # {팀: {공종: [사원명, ...]}}
        for name, data in self._current_emp_dict.items():
            team = data["team"]
            gj = data["gongjong"]
            if team not in team_gongjong:
                team_gongjong[team] = {}
            if gj not in team_gongjong[team]:
                team_gongjong[team][gj] = []
            team_gongjong[team][gj].append(name)

        for team in sorted(team_gongjong.keys()):
            gj_dict = team_gongjong[team]
            team_total = sum(len(v) for v in gj_dict.values())
            team_node = self.tree_team.insert("", tk.END, text=f"📂 {team}",
                                               values=(f"{team_total}명",), open=False)
            for gj in sorted(gj_dict.keys()):
                count = len(gj_dict[gj])
                self.tree_team.insert(team_node, tk.END, text=f"  🔧 {gj}",
                                       values=(f"{count}명",),
                                       tags=(f"{team}||{gj}",))

        # 오른쪽 목록 초기화
        for item in self.tree_names.get_children():
            self.tree_names.delete(item)
        self.lbl_name_title.config(text="👤 공종을 선택하세요")

        total = len(self._current_emp_dict)
        self.lbl_emp_title.config(
            text=f"👤 [{display_label}] 접속자 명단 — {total}명"
        )

    def _on_team_tree_select(self, event=None):
        """💡 왼쪽 팀/공종 트리에서 항목 클릭 → 오른쪽에 사원명 표시"""
        selected = self.tree_team.selection()
        if not selected:
            return

        item = selected[0]
        tags = self.tree_team.item(item, "tags")

        # 오른쪽 목록 초기화
        for row in self.tree_names.get_children():
            self.tree_names.delete(row)

        names_to_show = []

        if tags and "||" in tags[0]:
            # 공종 노드 클릭
            team, gj = tags[0].split("||", 1)
            for name, data in self._current_emp_dict.items():
                if data["team"] == team and data["gongjong"] == gj:
                    names_to_show.append(name)
            self.lbl_name_title.config(text=f"👤 [{team}] {gj} — {len(names_to_show)}명")
        else:
            # 팀 노드 클릭 → 해당 팀 전체
            item_text = self.tree_team.item(item, "text").replace("📂 ", "").strip()
            for name, data in self._current_emp_dict.items():
                if data["team"] == item_text:
                    names_to_show.append(name)
            self.lbl_name_title.config(text=f"👤 [{item_text}] 전체 — {len(names_to_show)}명")

        names_to_show.sort()
        for name in names_to_show:
            self.tree_names.insert("", tk.END, values=(name,))

    def hide_emp_list(self):
        # PanedWindow에서 명단 패널 제거
        if self._emp_panel_visible:
            self.frame_content.forget(self.frame_emp_area)
            self._emp_panel_visible = False
        self.clicked_zone = None
        self.clicked_group_zones = None
        self._current_emp_dict = {}

        # 💡 포커스 해제 → 모든 바 다시 보이게
        self._focused_bar_idx = None
        self._hovered_bar_idx = -1
        self._update_bar_highlight()

        # 💡 명단 닫으면 최소 크기/창 크기 원래대로 복원
        self.root.minsize(950, 700)
        current_w = self.root.winfo_width()
        self.root.geometry(f"{current_w}x700")

    # ==========================================
    # 🌐 서버 연동
    # ==========================================
    def auto_connect_and_start(self):
        self.lbl_status.config(text="⬇️ 서버에서 사원명부 연동 중...", fg=self.colors["primary"])
        self.root.update()

        try:
            headers = {"X-API-Key": API_KEY}
            res = requests.get(f"{SERVER_URL}/api/employees", headers=headers,
                               timeout=5, proxies={"http": None, "https": None})

            if res.status_code == 200:
                with open(TEMP_EMP_FILE, "wb") as f:
                    f.write(res.content)
                self.parse_employee_file(TEMP_EMP_FILE)
                self.fetch_zone_limits()  # 💡 서버에서 접속 제한 인원 받아오기
                self.debug_log(f"서버 연동 완료: IP {len(self.emp_by_ip)}건, PC {len(self.emp_by_pc)}건, 사원정보 {len(self.emp_info)}건")
                self.lbl_status.config(text="🟢 서버 연동 완료 (실시간 관제 중)", fg=self.colors["status_online"])

                # 💡 설정에 따라 기본 화면 결정
                default_region = self.settings.get("default_region", "__favorites__")
                if default_region == "__favorites__":
                    self.show_favorites()
                else:
                    default_label = self._get_button_label(default_region)
                    self.on_region_click(default_label)

                self.manual_refresh()
            elif res.status_code == 401:
                self.lbl_status.config(text="⚠️ 인증 실패: API KEY를 확인하세요.", fg=self.colors["status_offline"])
            else:
                self.lbl_status.config(text="⚠️ '사원명부.xlsx'가 서버에 없습니다.", fg=self.colors["status_offline"])

        except requests.exceptions.RequestException as e:
            self.lbl_status.config(text=f"❌ 연결 실패: {str(e)[:30]}...", fg=self.colors["status_offline"])

    def fetch_zone_limits(self):
        """💡 서버에서 접속 제한 인원 받아오기"""
        try:
            headers = {"X-API-Key": API_KEY}
            res = requests.get(f"{SERVER_URL}/api/zone_limits", headers=headers,
                               timeout=3, proxies={"http": None, "https": None})
            if res.status_code == 200:
                self.server_zone_limits = res.json()
                self.debug_log(f"접속 제한 수신: {len(self.server_zone_limits)}개 구역")
        except:
            pass  # 실패해도 기존 값 유지

    def parse_employee_file(self, filepath):
        app = None
        try:
            app = xw.App(visible=False)
            wb = app.books.open(filepath)
            sheet = wb.sheets[0]
            data = sheet.used_range.value

            if not data: return
            if not isinstance(data[0], list): data = [data]

            for row in data:
                if len(row) < 3: continue
                name = str(row[0]).strip() if row[0] is not None else ""
                ip = str(row[1]).strip() if row[1] is not None else ""
                pc_name = str(row[2]).strip() if row[2] is not None else ""
                team = str(row[3]).strip() if len(row) > 3 and row[3] is not None else "미지정"
                gongjong = str(row[4]).strip() if len(row) > 4 and row[4] is not None else "미지정"

                if name == "이름" or ip.lower() == "ip": continue
                if ip: self.emp_by_ip[ip] = name
                if pc_name: self.emp_by_pc[pc_name] = name
                # 💡 팀/공종/IP/PC 정보 저장
                self.emp_info[name] = {"team": team, "gongjong": gongjong, "ip": ip, "pc": pc_name}
        except Exception as e:
            self.debug_log(f"명부 파싱 에러: {e}")
        finally:
            if app:
                try: wb.close()
                except: pass
                app.quit()
            if os.path.exists(filepath):
                try: os.remove(filepath)
                except: pass

    def manual_refresh(self):
        """💡 수동 갱신 버튼 또는 자동 갱신 타이머에서 호출"""
        self.debug_log("갱신 시작")
        try:
            headers = {"X-API-Key": API_KEY}
            res = requests.get(f"{SERVER_URL}/api/logs", headers=headers,
                               timeout=5, proxies={"http": None, "https": None})

            if res.status_code == 200:
                with open(TEMP_LOG_FILE, "wb") as f:
                    f.write(res.content)

                self.parse_log_file(TEMP_LOG_FILE)
                self.fetch_zone_limits()
                self.refresh_chart()

                # 접속자 명단이 열려 있으면 함께 갱신
                if self.clicked_group_zones:
                    zones = self.clicked_group_zones
                    display_label = get_display_name(zones[0]) if len(zones) == 1 else "선택된 그룹"
                    current_text = self.lbl_emp_title.cget("text")
                    bracket_match = re.search(r'\[(.*?)\]', current_text)
                    if bracket_match:
                        display_label = bracket_match.group(1)
                    self.show_emp_list_for_zones(zones, display_label)

                current_time = time.strftime('%H:%M:%S')
                self.lbl_time.config(text=f"마지막 갱신: {current_time}", fg="#757575")
            elif res.status_code == 401:
                self.lbl_time.config(text="⚠️ 인증 실패", fg=self.colors["status_offline"])
            else:
                self.lbl_time.config(text="오늘자 로그가 아직 없습니다.", fg="gray")

        except Exception as e:
            self.lbl_time.config(text=f"⚠️ 통신 불안정: {str(e)[:20]}...", fg=self.colors["status_offline"])

        # 💡 자동 갱신 주기에 따라 다음 타이머 예약
        self._schedule_auto_refresh()

    def _on_refresh_interval_change(self, event=None):
        """갱신 주기 콤보박스 변경 시"""
        # 설정 저장
        self.settings["refresh_interval"] = self._refresh_interval_var.get()
        save_settings(self.settings)
        # 기존 타이머 취소 후 새 주기로 예약
        if self._update_job:
            self.root.after_cancel(self._update_job)
            self._update_job = None
        self._schedule_auto_refresh()

    def _schedule_auto_refresh(self):
        """현재 선택된 주기에 따라 자동 갱신 타이머 예약"""
        interval_map = {"수동": 0, "1분": 60000, "5분": 300000, "10분": 600000}
        interval = interval_map.get(self._refresh_interval_var.get(), 0)

        if interval > 0:
            self._update_job = self.root.after(interval, self.manual_refresh)

    # ==========================================
    # 📄 로그 파싱
    # ==========================================
    def parse_log_file(self, filepath):
        try:
            # 💡 인코딩 자동 감지: UTF-8 → cp949 → latin-1(절대 실패 안함)
            for encoding in ['utf-8-sig', 'cp949']:
                try:
                    with open(filepath, 'r', encoding=encoding) as f:
                        raw = f.read()
                    break
                except UnicodeDecodeError:
                    continue
            else:
                # 💡 최후 수단: latin-1은 모든 바이트를 읽을 수 있음
                with open(filepath, 'r', encoding='latin-1') as f:
                    raw = f.read()

            import io
            reader = csv.reader(io.StringIO(raw))
            header = next(reader)

            temp_active_users = {}
            temp_online_users = set()  # 💡 Client 데이터를 보낸 PC (켜져있는 PC)
            for row in reader:
                if len(row) < 6: continue
                log_ip = row[1].strip()
                log_pc = row[2].strip()
                process_name = row[3].strip()
                status = row[4]
                details = row[5]

                emp_name = self.emp_by_ip.get(log_ip)
                if not emp_name:
                    emp_name = self.emp_by_pc.get(log_pc)
                    if not emp_name:
                        emp_name = f"미등록({log_pc})"

                # 💡 데이터를 보낸 사원 = PC가 켜져 있는 상태
                temp_online_users.add(emp_name)

                key = (emp_name, process_name)

                if "미실행" in status or "X" in status:
                    if key in temp_active_users:
                        del temp_active_users[key]
                    continue

                zones_found = []
                proc_lower = process_name.lower()

                # 💡 S5D.exe — 창제목 패턴: ... - 구역명 - ...
                if proc_lower == "s5d.exe":
                    for window_title in details.split(' / '):
                        matches = re.findall(r' - (.*?) - ', window_title)
                        if matches:
                            zone_name = matches[0].strip()
                            zones_found.append(zone_name)
                        else:
                            parts = window_title.split(' - ')
                            if len(parts) >= 3:
                                zone_name = parts[1].strip()
                                zones_found.append(zone_name)
                    # 💡 구역 추출 실패 → 집계 제외 (초기화 대기, 비작업 창 등)
                    if not zones_found:
                        continue

                # 💡 DDWORKS — 창제목 패턴: ... [구역명] 또는 [구역명_접미사]
                elif proc_lower == "dinno.hu3d.wpf.hookupdesigner.exe":
                    matches = re.findall(r'\[(.*?)\]', details)
                    for match in matches:
                        if "_" in match:
                            zone_name = match.rsplit('_', 1)[0].strip()
                        else:
                            zone_name = match.strip()
                        zones_found.append(zone_name)
                    # 💡 Hookup Designer 창이 안 열린 상태 → 집계 제외
                    if not zones_found:
                        continue
                else:
                    continue  # 💡 알 수 없는 프로세스도 집계 제외

                temp_active_users[key] = zones_found

            self.active_users = temp_active_users
            self.online_users = temp_online_users
            self.debug_log(f"파싱 완료: active_users {len(self.active_users)}건, 온라인 PC {len(self.online_users)}대")
        except Exception as e:
            self.debug_log(f"파싱 오류: {e}")
            self.lbl_time.config(text=f"⚠️ 파싱 오류: {str(e)[:40]}", fg=self.colors["status_offline"])
        finally:
            if os.path.exists(filepath):
                try: os.remove(filepath)
                except: pass

    # ==========================================
    # 🐛 디버그 창
    # ==========================================
    def open_debug_window(self):
        dlg = tk.Toplevel(self.root)
        dlg.title("🔍 뷰어 디버그 상태창")
        dlg.geometry("850x650")
        dlg.attributes("-topmost", True)
        dlg.configure(bg="#1E1E1E")

        stats_frame = tk.Frame(dlg, bg="#333333", pady=8, padx=15)
        stats_frame.pack(fill=tk.X)

        lbl_server = tk.Label(stats_frame, text=f"🌐 접속 서버: {SERVER_URL}", font=("Consolas", 10), fg="#A6E22E",
                              bg="#333333")
        lbl_server.pack(side=tk.LEFT)

        def get_total_counts():
            total_unique_users = len(set(emp_name for emp_name, _ in self.active_users.keys()))
            total_processes = len(self.active_users)
            return total_unique_users, total_processes

        u_count, p_count = get_total_counts()
        lbl_stats = tk.Label(stats_frame, text=f"👥 현재 접속: {u_count}명 (프로세스 {p_count}개)", font=("맑은 고딕", 11, "bold"),
                             fg="#FD971F", bg="#333333")
        lbl_stats.pack(side=tk.RIGHT)

        notebook = ttk.Notebook(dlg)
        notebook.pack(fill=tk.BOTH, expand=True, padx=5, pady=5)

        tab_logs = tk.Frame(notebook, bg="#1E1E1E")
        notebook.add(tab_logs, text=" 📝 실시간 로그 ")

        tab_users = tk.Frame(notebook, bg="#1E1E1E")
        notebook.add(tab_users, text=" 👥 팀별 접속자 현황 (미실행 포함) ")

        txt = scrolledtext.ScrolledText(tab_logs, font=("Consolas", 9), bg="#1E1E1E", fg="#D4D4D4", wrap=tk.WORD)
        txt.pack(fill=tk.BOTH, expand=True, padx=2, pady=2)

        for line in self.debug_logs:
            txt.insert(tk.END, line + "\n")
        txt.see(tk.END)
        txt.config(state=tk.DISABLED)

        txt_users = scrolledtext.ScrolledText(tab_users, font=("맑은 고딕", 10), bg="#1E1E1E", fg="#D4D4D4", wrap=tk.WORD)
        txt_users.pack(fill=tk.BOTH, expand=True, padx=2, pady=2)

        last_user_state_str = [""]
        last_log_count = [len(self.debug_logs)]

        def refresh_debug_ui():
            if not dlg.winfo_exists():
                return

            current_u, current_p = get_total_counts()
            lbl_stats.config(text=f"👥 현재 접속: {current_u}명 (프로세스 {current_p}개)")

            # 로그 탭 갱신
            current_len = len(self.debug_logs)
            if current_len > last_log_count[0]:
                txt.config(state=tk.NORMAL)
                for line in self.debug_logs[last_log_count[0]:]:
                    txt.insert(tk.END, line + "\n")
                txt.see(tk.END)
                txt.config(state=tk.DISABLED)
                last_log_count[0] = current_len
            elif current_len < last_log_count[0]:
                txt.config(state=tk.NORMAL)
                txt.delete(1.0, tk.END)
                for line in self.debug_logs:
                    txt.insert(tk.END, line + "\n")
                txt.see(tk.END)
                txt.config(state=tk.DISABLED)
                last_log_count[0] = current_len

            # 팀별 접속자 현황 갱신
            user_locs = {}
            for (emp_name, proc_name), zones in self.active_users.items():
                if emp_name not in user_locs:
                    user_locs[emp_name] = set()

                for z in zones:
                    specific_key = f"{z}:{proc_name.lower()}"
                    if specific_key in ZONE_DISPLAY_MAP:
                        disp_name = ZONE_DISPLAY_MAP[specific_key]
                    else:
                        disp_name = get_display_name(z)

                    user_locs[emp_name].add(disp_name)

            all_users = set(self.emp_info.keys()).union(set(user_locs.keys()))

            user_details = []
            for emp_name in all_users:
                info = self.emp_info.get(emp_name, {"team": "미지정", "gongjong": "미지정", "ip": "IP없음", "pc": "PC없음"})
                team = info.get("team", "미지정")
                gj = info.get("gongjong", "미지정")
                ip_str = info.get("ip", "IP없음")
                pc_str = info.get("pc", "PC없음")

                if emp_name in user_locs:
                    # S5D/DDWORKS 실행 중 + 구역 정보 있음
                    zones_str = ", ".join(sorted(user_locs[emp_name]))
                    status_mark = "🟢"
                    loc_str = f"[{zones_str}]"
                elif emp_name in self.online_users:
                    # PC 켜져 있지만 S5D/DDWORKS 미실행
                    status_mark = "🟢"
                    loc_str = "S5D,DDWORKS 실행안됨"
                else:
                    # PC 꺼져 있음 (Client 데이터 없음)
                    status_mark = "❌"
                    loc_str = "(PC OFF)"

                user_details.append((team, gj, emp_name, ip_str, pc_str, status_mark, loc_str))

            user_details.sort(key=lambda x: (x[0], x[1], x[2]))

            display_lines = []
            current_team = None
            for team, gj, emp_name, ip_str, pc_str, status_mark, loc_str in user_details:
                if current_team != team:
                    if current_team is not None:
                        display_lines.append("")
                    display_lines.append(f"🏢 ━━━ {team} ━━━")
                    current_team = team

                display_lines.append(f" {status_mark} {emp_name}/{gj}  |  PC: {pc_str}  |  IP: {ip_str}  |  {loc_str}")

            current_state_str = "\n".join(display_lines)

            if current_state_str != last_user_state_str[0]:
                txt_users.config(state=tk.NORMAL)
                txt_users.delete(1.0, tk.END)
                txt_users.insert(tk.END, current_state_str)
                txt_users.config(state=tk.DISABLED)
                last_user_state_str[0] = current_state_str

            dlg.after(1000, refresh_debug_ui)

        dlg.after(100, refresh_debug_ui)

    # ==========================================
    # 📥 현황 내려받기 (Excel)
    # ==========================================
    def _collect_zone_data(self, zone_list):
        """구역 리스트 → {구역표시명: [(팀/공종, 이름), ...]} 데이터 수집"""
        result = {}
        for zone_key in zone_list:
            display = get_display_name(zone_key)
            users = self._count_zone_users(zone_key)
            entries = []
            for emp_name in sorted(users):
                info = self.emp_info.get(emp_name, {"team": "미지정", "gongjong": "미지정"})
                team_gj = f"{info['team']}/{info['gongjong']}"
                entries.append((team_gj, emp_name))
            # 팀/공종 기준 정렬
            entries.sort(key=lambda x: (x[0], x[1]))
            result[display] = entries
        return result

    def _get_all_region_zones(self):
        """모든 지역의 구역 키를 지역명과 함께 반환: [(지역표시명, [zone_key, ...])]"""
        result = []
        for btn_label, internal_region in REGION_BUTTONS.items():
            config = REGION_CONFIG.get(internal_region, {})
            zones = []
            if config.get("grouped"):
                for sub_zones in config.get("groups", {}).values():
                    zones.extend(sub_zones)
            else:
                zones = config.get("zones", [])
            if zones:
                result.append((btn_label, zones))
        return result

    def export_to_excel(self):
        """💡 현황 데이터를 Excel 파일로 내려받기"""
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
        except ImportError:
            from tkinter import messagebox
            messagebox.showerror("모듈 필요", "openpyxl이 설치되어 있지 않습니다.\npip install openpyxl 로 설치해주세요.")
            return

        if not self.active_users:
            from tkinter import messagebox
            messagebox.showinfo("알림", "집계된 데이터가 없습니다.\n먼저 🔄 갱신을 실행해주세요.")
            return

        # 💡 임시 파일에 저장 후 엑셀로 바로 열기
        import tempfile
        now_str = datetime.now().strftime("%Y%m%d_%H%M")
        temp_dir = tempfile.gettempdir()
        filepath = os.path.join(temp_dir, f"접속현황_{now_str}.xlsx")

        # 스타일 정의
        header_font = Font(name="맑은 고딕", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="00479A", end_color="00479A", fill_type="solid")
        zone_font = Font(name="맑은 고딕", size=11, bold=True, color="00479A")
        zone_fill = PatternFill(start_color="E8EAF6", end_color="E8EAF6", fill_type="solid")
        data_font = Font(name="맑은 고딕", size=10)
        center_align = Alignment(horizontal="center", vertical="center")
        left_align = Alignment(horizontal="left", vertical="center")
        thin_border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )

        def write_zone_data(ws, zone_data, start_row=1):
            """시트에 구역별 데이터 작성, 마지막 행 번호 반환"""
            row = start_row
            for zone_name, entries in zone_data.items():
                # 구역명 헤더
                cell = ws.cell(row=row, column=1, value=f"📍 {zone_name}")
                cell.font = zone_font
                cell.fill = zone_fill
                cell.alignment = left_align
                cell.border = thin_border
                ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
                ws.cell(row=row, column=2).border = thin_border

                count_cell = ws.cell(row=row, column=3, value=f"{len(entries)}명")
                count_cell.font = zone_font
                count_cell.fill = zone_fill
                count_cell.alignment = center_align
                count_cell.border = thin_border
                row += 1

                # 칼럼 헤더
                for col, text in [(1, "팀/공종"), (2, "이름")]:
                    cell = ws.cell(row=row, column=col, value=text)
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = center_align
                    cell.border = thin_border
                row += 1

                if entries:
                    for team_gj, name in entries:
                        ws.cell(row=row, column=1, value=team_gj).font = data_font
                        ws.cell(row=row, column=1).alignment = left_align
                        ws.cell(row=row, column=1).border = thin_border
                        ws.cell(row=row, column=2, value=name).font = data_font
                        ws.cell(row=row, column=2).alignment = center_align
                        ws.cell(row=row, column=2).border = thin_border
                        row += 1
                else:
                    ws.cell(row=row, column=1, value="접속자 없음").font = data_font
                    ws.cell(row=row, column=1).alignment = center_align
                    ws.cell(row=row, column=1).border = thin_border
                    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
                    ws.cell(row=row, column=2).border = thin_border
                    row += 1

                row += 1  # 구역 간 빈 행
            return row

        try:
            wb = Workbook()

            # ======== 시트 1: 즐겨찾기 ========
            ws1 = wb.active
            ws1.title = "즐겨찾기"
            ws1.column_dimensions['A'].width = 22
            ws1.column_dimensions['B'].width = 18
            ws1.column_dimensions['C'].width = 10

            fav_zones = self.settings.get("favorite_zones", [])
            if fav_zones:
                fav_data = self._collect_zone_data(fav_zones)
                write_zone_data(ws1, fav_data)
            else:
                ws1.cell(row=1, column=1, value="즐겨찾기 구역이 설정되지 않았습니다.").font = data_font

            # ======== 시트 2: 평기화천해 ========
            ws2 = wb.create_sheet(title="평기화천해")
            ws2.column_dimensions['A'].width = 22
            ws2.column_dimensions['B'].width = 18
            ws2.column_dimensions['C'].width = 10

            row = 1
            for region_label, zone_keys in self._get_all_region_zones():
                # 지역 대제목
                cell = ws2.cell(row=row, column=1, value=f"■ {region_label}")
                cell.font = Font(name="맑은 고딕", size=12, bold=True, color="003C82")
                ws2.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
                row += 1

                region_data = self._collect_zone_data(zone_keys)
                row = write_zone_data(ws2, region_data, start_row=row)
                row += 1  # 지역 간 추가 간격

            wb.save(filepath)

            # 💡 저장된 임시 파일을 엑셀로 바로 열기 (사용자가 다른이름저장 가능)
            os.startfile(filepath)

        except Exception as e:
            from tkinter import messagebox
            messagebox.showerror("저장 오류", f"파일 저장 중 오류가 발생했습니다.\n{str(e)}")

    # ==========================================
    # 🚪 종료 시 정리
    # ==========================================
    def on_close(self):
        """💡 창의 X 버튼 → 트레이로 숨김 (백그라운드 모니터링은 계속).
        트레이가 없으면(예: 트레이 생성 실패) 곧바로 완전 종료한다."""
        if self.tray_icon is not None and not self._quitting:
            self.hide_to_tray()
        else:
            self.real_quit()


# ==========================================================
# 🚀 통합 실행 진입점
# ==========================================================
def main():
    # 💡 트레이 모드(부팅 자동 실행)로 시작했는지 확인
    start_minimized = (TRAY_START_FLAG in sys.argv)

    # 💡 작업 디렉터리를 앱 폴더로 고정 (설정/로그/임시파일 경로 안정화)
    try:
        os.chdir(APP_DIR)
    except Exception:
        pass

    # 💡 1) 중복 실행 방지
    if not acquire_single_instance():
        sys.exit(0)

    # 💡 2) Windows 부팅 시 자동 실행(트레이 모드) 등록
    register_autostart()

    # 💡 3) 뷰어 GUI 생성
    root = tk.Tk()
    viewer = MonitorViewer(root)

    # 💡 4) 백그라운드 모니터링 에이전트 시작 (로그는 디버그 창으로도 전달)
    agent = MonitoringAgent(log_callback=viewer.add_log_line)
    viewer.attach_agent(agent)
    agent.start()

    # 💡 5) 시스템 트레이 아이콘을 별도 스레드에서 실행
    #     (tkinter는 메인 스레드, pystray는 워커 스레드 — Windows에서 안전)
    tray_icon = None
    try:
        tray_icon = build_tray_icon(viewer)
        viewer.attach_tray(tray_icon)
        threading.Thread(target=tray_icon.run, daemon=True).start()
    except Exception as e:
        print(f"[트레이 생성 실패 — 트레이 없이 실행] {e}")

    # 💡 6) 트레이 모드로 시작했다면 창을 숨긴 채 백그라운드 상주
    if start_minimized and tray_icon is not None:
        root.after(100, root.withdraw)

    # 💡 7) 메인 루프
    try:
        root.mainloop()
    except KeyboardInterrupt:
        viewer.real_quit()
    finally:
        agent.stop()
        if tray_icon is not None:
            try:
                tray_icon.stop()
            except Exception:
                pass


if __name__ == "__main__":
    try:
        main()
    except Exception:
        logging.basicConfig(
            filename=os.path.join(APP_DIR, "monitor_error.log"),
            level=logging.ERROR,
        )
        logging.error("치명적 오류", exc_info=True)
        raise
